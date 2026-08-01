"""Phase 2: build professional classroom PPTX from structured deck + theme."""

from __future__ import annotations

from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import re

from app.services.lesson_planner.export.pptx_themes import PptTheme, resolve_theme

# Widescreen 16:9
_SLIDE_W = Inches(13.333)
_SLIDE_H = Inches(7.5)


def _rgb(rgb: tuple[int, int, int]) -> RGBColor:
    return RGBColor(rgb[0], rgb[1], rgb[2])


def _set_run(run, *, size: int, bold: bool = False, color: tuple[int, int, int], font: str) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    run.font.name = font


def _fill_solid(shape, rgb: tuple[int, int, int]) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(rgb)
    shape.line.fill.background()


def _add_rect(slide, left, top, width, height, fill: tuple[int, int, int]):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    _fill_solid(shape, fill)
    return shape


def _add_round_rect(slide, left, top, width, height, fill: tuple[int, int, int]):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    _fill_solid(shape, fill)
    return shape


def _textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def _write_paragraphs(
    tf,
    lines: list[str],
    *,
    size: int,
    color: tuple[int, int, int],
    font: str,
    bold: bool = False,
    bullet: bool = False,
    space_after: int = 8,
) -> None:
    tf.clear()
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space_after)
        if bullet:
            p.level = 0
        run = p.add_run()
        run.text = (("• " if bullet else "") + line) if line else ""
        _set_run(run, size=size, bold=bold, color=color, font=font)


def _paint_background(slide, theme: PptTheme) -> None:
    _add_rect(slide, 0, 0, _SLIDE_W, _SLIDE_H, theme.bg)
    # Left accent bar
    _add_rect(slide, 0, 0, Inches(0.18), _SLIDE_H, theme.primary)


def _add_notes(slide, notes: str) -> None:
    if not notes:
        return
    slide.notes_slide.notes_text_frame.text = notes


_ICON_LABELS = {
    "engage": "ENGAGE",
    "example": "EXAMPLE",
    "try": "TRY THIS",
    "check": "CHECK",
    "diagram": "DIAGRAM",
    "formula": "FORMULA",
    "remember": "REMEMBER",
}


def _icon_chip(slide, icon: str, theme: PptTheme, *, left=Inches(0.7), top=Inches(0.32)) -> None:
    label = _ICON_LABELS.get((icon or "").lower())
    if not label:
        return
    width = Inches(1.55)
    _add_round_rect(slide, left, top, width, Inches(0.38), theme.primary)
    box = _textbox(slide, left, top + Inches(0.02), width, Inches(0.34))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    _set_run(run, size=10, bold=True, color=theme.on_primary, font=theme.font_body)


def _header(slide, item: dict[str, Any], theme: PptTheme) -> None:
    icon = (item.get("icon") or "").strip()
    side = (item.get("side_heading") or "").strip()
    title_top = Inches(0.7)
    if icon:
        _icon_chip(slide, icon, theme)
        if side:
            chip = _textbox(slide, Inches(2.4), Inches(0.35), Inches(9.8), Inches(0.4))
            _write_paragraphs(
                chip.text_frame,
                [side.upper()],
                size=12,
                color=theme.accent,
                font=theme.font_body,
                bold=True,
            )
        title_top = Inches(0.85)
    elif side:
        chip = _textbox(slide, Inches(0.7), Inches(0.35), Inches(11.5), Inches(0.4))
        _write_paragraphs(
            chip.text_frame,
            [side.upper()],
            size=12,
            color=theme.accent,
            font=theme.font_body,
            bold=True,
        )
    title = _textbox(slide, Inches(0.7), title_top, Inches(11.5), Inches(1.0))
    _write_paragraphs(
        title.text_frame,
        [item.get("title") or "Slide"],
        size=28,
        color=theme.text,
        font=theme.font_title,
        bold=True,
    )


def _callout(slide, text: str, theme: PptTheme, top=Inches(6.2)) -> None:
    if not text:
        return
    cleaned = re.sub(r"[*_`]+", "", str(text)).strip()
    if not cleaned:
        return
    if cleaned.lower().startswith("remember:"):
        cleaned = cleaned.split(":", 1)[1].strip()
    _add_round_rect(slide, Inches(0.7), top, Inches(11.8), Inches(0.85), theme.surface)
    _add_rect(slide, Inches(0.7), top, Inches(0.12), Inches(0.85), theme.accent)
    box = _textbox(slide, Inches(1.05), top + Inches(0.15), Inches(11.2), Inches(0.6))
    _write_paragraphs(
        box.text_frame,
        [f"Remember: {cleaned}"],
        size=14,
        color=theme.text,
        font=theme.font_body,
        bold=True,
    )


def _render_title(slide, item: dict[str, Any], theme: PptTheme, meta: dict[str, str]) -> None:
    _add_rect(slide, 0, 0, _SLIDE_W, _SLIDE_H, theme.primary)
    _add_rect(slide, 0, Inches(5.4), _SLIDE_W, Inches(2.1), theme.accent)

    eyebrow = _textbox(slide, Inches(0.9), Inches(1.6), Inches(11), Inches(0.5))
    _write_paragraphs(
        eyebrow.text_frame,
        [meta.get("eyebrow") or "Classroom Lesson"],
        size=16,
        color=theme.on_primary,
        font=theme.font_body,
        bold=True,
    )

    title_box = _textbox(slide, Inches(0.9), Inches(2.2), Inches(11.2), Inches(2.2))
    _write_paragraphs(
        title_box.text_frame,
        [item.get("title") or meta.get("chapter") or "Lesson"],
        size=40,
        color=theme.on_primary,
        font=theme.font_title,
        bold=True,
        space_after=6,
    )

    sub = f"{meta.get('subject', '')}  ·  {meta.get('grade', '')}".strip(" ·")
    if sub:
        sub_box = _textbox(slide, Inches(0.9), Inches(5.7), Inches(11), Inches(0.6))
        _write_paragraphs(
            sub_box.text_frame,
            [sub],
            size=18,
            color=theme.on_primary,
            font=theme.font_body,
        )


def _render_section(slide, item: dict[str, Any], theme: PptTheme) -> None:
    _paint_background(slide, theme)
    _add_round_rect(slide, Inches(1.2), Inches(2.4), Inches(10.8), Inches(2.4), theme.surface)
    box = _textbox(slide, Inches(1.6), Inches(2.8), Inches(10), Inches(1.6))
    _write_paragraphs(
        box.text_frame,
        [item.get("title") or "Section"],
        size=36,
        color=theme.primary,
        font=theme.font_title,
        bold=True,
    )


def _render_figure(slide, item: dict[str, Any], theme: PptTheme) -> None:
    """Image + teaching points — uses textbook figure when figure_path is set."""
    import os

    _paint_background(slide, theme)
    _header(slide, item, theme)
    _add_round_rect(slide, Inches(0.7), Inches(1.85), Inches(6.0), Inches(4.5), theme.surface)
    _add_round_rect(slide, Inches(7.0), Inches(1.85), Inches(5.5), Inches(4.5), theme.surface)

    left_body = _textbox(slide, Inches(1.0), Inches(2.1), Inches(5.4), Inches(3.9))
    points = item.get("bullets") or []
    caption = (item.get("figure_caption") or "").strip()
    lines = points or ([caption] if caption else ["Look carefully at the diagram."])
    _write_paragraphs(
        left_body.text_frame,
        lines,
        size=16,
        color=theme.text,
        font=theme.font_body,
        bullet=True,
        space_after=10,
    )

    path = str(item.get("figure_path") or "")
    if path and os.path.isfile(path):
        try:
            # Leave margin inside the right card.
            slide.shapes.add_picture(path, Inches(7.25), Inches(2.15), width=Inches(5.0))
        except Exception:
            _figure_placeholder(slide, theme, caption or "Textbook figure")
    else:
        _figure_placeholder(slide, theme, caption or "Textbook figure")


def _figure_placeholder(slide, theme: PptTheme, caption: str) -> None:
    box = _textbox(slide, Inches(7.3), Inches(3.2), Inches(4.9), Inches(1.8))
    _write_paragraphs(
        box.text_frame,
        ["[ Diagram ]", caption[:90]],
        size=14,
        color=theme.muted,
        font=theme.font_body,
        bold=True,
        space_after=8,
    )


def _render_bullets(slide, item: dict[str, Any], theme: PptTheme) -> None:
    _paint_background(slide, theme)
    _header(slide, item, theme)
    card = _add_round_rect(slide, Inches(0.7), Inches(1.9), Inches(11.8), Inches(4.0), theme.surface)
    _ = card
    body = _textbox(slide, Inches(1.05), Inches(2.15), Inches(11.1), Inches(3.5))
    _write_paragraphs(
        body.text_frame,
        item.get("bullets") or [],
        size=20,
        color=theme.text,
        font=theme.font_body,
        bullet=True,
        space_after=12,
    )
    _callout(slide, item.get("callout") or "", theme)


def _render_two_column(slide, item: dict[str, Any], theme: PptTheme) -> None:
    _paint_background(slide, theme)
    _header(slide, item, theme)
    left = item.get("bullets") or []
    right = item.get("right_bullets") or []
    _add_round_rect(slide, Inches(0.7), Inches(1.9), Inches(5.7), Inches(4.0), theme.surface)
    _add_round_rect(slide, Inches(6.8), Inches(1.9), Inches(5.7), Inches(4.0), theme.surface)

    left_label = _textbox(slide, Inches(1.0), Inches(2.1), Inches(5.1), Inches(0.4))
    _write_paragraphs(left_label.text_frame, ["Concept"], size=13, color=theme.accent, font=theme.font_body, bold=True)
    left_body = _textbox(slide, Inches(1.0), Inches(2.55), Inches(5.1), Inches(3.1))
    _write_paragraphs(left_body.text_frame, left, size=16, color=theme.text, font=theme.font_body, bullet=True, space_after=10)

    right_label = _textbox(slide, Inches(7.1), Inches(2.1), Inches(5.1), Inches(0.4))
    _write_paragraphs(right_label.text_frame, ["Example"], size=13, color=theme.accent, font=theme.font_body, bold=True)
    right_body = _textbox(slide, Inches(7.1), Inches(2.55), Inches(5.1), Inches(3.1))
    _write_paragraphs(right_body.text_frame, right or left[len(left) // 2 :], size=16, color=theme.text, font=theme.font_body, bullet=True, space_after=10)


def _render_steps(slide, item: dict[str, Any], theme: PptTheme) -> None:
    steps = [s for s in (item.get("bullets") or []) if str(s).strip()][:4]
    # Fewer than 2 steps looks broken as cards — fall back to bullets.
    if len(steps) < 2:
        _render_bullets(slide, {**item, "bullets": steps or item.get("bullets") or []}, theme)
        return
    _paint_background(slide, theme)
    _header(slide, item, theme)
    width = Inches(2.7)
    gap = Inches(0.25)
    start = Inches(0.7)
    for i, step in enumerate(steps):
        left = start + i * (width + gap)
        _add_round_rect(slide, left, Inches(2.1), width, Inches(3.6), theme.surface)
        badge = _add_round_rect(slide, left + Inches(0.9), Inches(2.35), Inches(0.9), Inches(0.9), theme.primary)
        nbox = _textbox(slide, left + Inches(0.9), Inches(2.45), Inches(0.9), Inches(0.7))
        tf = nbox.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = str(i + 1)
        _set_run(run, size=22, bold=True, color=theme.on_primary, font=theme.font_title)
        text = step.lstrip("0123456789.-) ").strip()
        tbox = _textbox(slide, left + Inches(0.2), Inches(3.5), width - Inches(0.4), Inches(2.0))
        _write_paragraphs(tbox.text_frame, [text], size=14, color=theme.text, font=theme.font_body, bold=False, space_after=4)
        _ = badge


def _render_big_idea(slide, item: dict[str, Any], theme: PptTheme) -> None:
    _paint_background(slide, theme)
    _header(slide, item, theme)
    idea = next(
        (str(b).strip() for b in (item.get("bullets") or []) if str(b).strip()),
        str(item.get("callout") or item.get("title") or ""),
    )
    _add_round_rect(slide, Inches(1.2), Inches(2.3), Inches(10.8), Inches(3.2), theme.surface)
    _add_rect(slide, Inches(1.2), Inches(2.3), Inches(0.18), Inches(3.2), theme.primary)
    box = _textbox(slide, Inches(1.8), Inches(2.9), Inches(9.8), Inches(2.2))
    _write_paragraphs(
        box.text_frame,
        [idea],
        size=28,
        color=theme.text,
        font=theme.font_title,
        bold=True,
    )
    _callout(slide, item.get("callout") or "Say it in your own words.", theme)


def _render_summary(slide, item: dict[str, Any], theme: PptTheme) -> None:
    _paint_background(slide, theme)
    _header(slide, item, theme)
    _add_round_rect(slide, Inches(0.7), Inches(1.9), Inches(11.8), Inches(4.5), theme.surface)
    body = _textbox(slide, Inches(1.1), Inches(2.2), Inches(11.0), Inches(3.8))
    title = str(item.get("title") or "").strip().lower()
    lines = []
    for b in item.get("bullets") or []:
        text = str(b).strip().lstrip("✓ ").strip()
        if not text or text.lower() in {title, f"✓ {title}"}:
            continue
        lines.append(f"✓  {text}")
    if not lines:
        lines = ["✓  Review today's key ideas", "✓  Ask one doubt before you leave"]
    _write_paragraphs(
        body.text_frame,
        lines,
        size=20,
        color=theme.text,
        font=theme.font_body,
        space_after=14,
    )


_LAYOUT_RENDERERS = {
    "title": _render_title,
    "section": _render_section,
    "bullets": _render_bullets,
    "two_column": _render_two_column,
    "steps": _render_steps,
    "big_idea": _render_big_idea,
    "summary": _render_summary,
    "figure": _render_figure,
}


def build_classroom_pptx(
    path,
    *,
    slides: list[dict[str, Any]],
    theme_id: str | None,
    meta: dict[str, str],
) -> None:
    theme = resolve_theme(theme_id)
    prs = Presentation()
    prs.slide_width = _SLIDE_W
    prs.slide_height = _SLIDE_H
    blank = prs.slide_layouts[6]  # blank

    deck = list(slides)
    if not deck:
        deck = [
            {
                "title": meta.get("chapter") or "Lesson",
                "layout": "title",
                "bullets": [],
                "side_heading": "",
                "callout": "",
                "speaker_notes": "",
            }
        ]

    for item in deck:
        slide = prs.slides.add_slide(blank)
        layout = item.get("layout") or "bullets"
        renderer = _LAYOUT_RENDERERS.get(layout, _render_bullets)
        if layout == "title":
            renderer(slide, item, theme, meta)
        else:
            renderer(slide, item, theme)
        _add_notes(slide, str(item.get("speaker_notes") or ""))

    prs.save(path)
