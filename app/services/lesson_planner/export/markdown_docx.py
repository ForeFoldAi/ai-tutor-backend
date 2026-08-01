from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt

from app.services.lesson_planner.export.markdown_pdf import ARTIFACT_TITLES


def _strip_md_bold(text: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"\1", text)


def _is_table_sep(line: str) -> bool:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.match(r"^:?-{2,}:?$", c) for c in cells if c)


def add_markdown_to_doc(doc: Document, markdown: str) -> None:
    for raw in markdown.replace("\r\n", "\n").split("\n"):
        stripped = raw.strip()
        if not stripped or stripped == "---":
            continue
        if stripped.startswith("|") and "|" in stripped[1:]:
            if _is_table_sep(stripped):
                continue
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            stripped = " | ".join(cells)
        if stripped.startswith("# "):
            doc.add_heading(_strip_md_bold(stripped[2:]), level=0)
        elif stripped.startswith("## "):
            doc.add_heading(_strip_md_bold(stripped[3:]), level=1)
        elif stripped.startswith("### "):
            doc.add_heading(_strip_md_bold(stripped[4:]), level=2)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(_strip_md_bold(stripped[2:]), style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
        elif re.match(r"^\d+\.\s", stripped):
            p = doc.add_paragraph(_strip_md_bold(stripped), style="List Number")
            p.paragraph_format.space_after = Pt(2)
        else:
            p = doc.add_paragraph(_strip_md_bold(stripped))
            p.paragraph_format.space_after = Pt(4)


def export_docx_document(
    path: Path,
    *,
    title: str,
    subtitle: str,
    sections: list[tuple[str, Any]],
) -> None:
    from app.services.lesson_planner.export.templates import render_artifact_lines

    doc = Document()
    doc.add_heading(title, 0)
    doc.add_paragraph(subtitle)

    for artifact_key, content in sections:
        section_title = ARTIFACT_TITLES.get(artifact_key, artifact_key.replace("_", " ").title())
        if len(sections) > 1:
            doc.add_heading(section_title, level=1)
        if isinstance(content, dict) and content.get("format") == "markdown" and content.get("markdown"):
            add_markdown_to_doc(doc, str(content["markdown"]))
        else:
            for line in render_artifact_lines(artifact_key, content):
                doc.add_paragraph(line)

    doc.save(path)
